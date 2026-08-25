"""Multimodal ingestion: Docling -> Bedrock (Embed v4 + VLM) -> records.

Trimmed for the Step Functions pipeline: the S3 Docling parse-cache and the VLM
description sidecar cache are removed (each Lambda invocation parses fresh), and
the Hatchet/Postgres coupling is gone. The record builders (text/figure/table/
chart/page) and the LanceDB write helpers are otherwise unchanged.

Handles PDF, PPTX and XLSX, producing one LanceDB table with a single unified
multimodal vector column so text chunks, figure descriptions, page screenshots
and tables all live in the same vector space.
"""

from __future__ import annotations

import hashlib
import io
import logging
import os
import subprocess
import tempfile
from collections import defaultdict
from concurrent.futures import ThreadPoolExecutor
from dataclasses import dataclass, field, replace
from pathlib import Path
from typing import Callable, Optional

import lancedb
from docling.chunking import HybridChunker
from docling.datamodel.base_models import InputFormat
from docling.datamodel.pipeline_options import (
    PaginatedPipelineOptions,
    PdfPipelineOptions,
)
from docling.document_converter import (
    DocumentConverter,
    ExcelFormatOption,
    PdfFormatOption,
    PowerpointFormatOption,
)
from docling_core.types.doc import PictureItem, TableItem

from common import SETTINGS, Bedrock, DocChunk, Settings

logger = logging.getLogger("semantic.ingest")

SUPPORTED_SUFFIXES = {".pdf", ".pptx", ".xlsx"}

# Chunker tokenizer, pinned to an immutable commit rather than a mutable branch
# so a compromised or retagged upstream repo can't swap the model out from under
# us. Must stay in sync with the bake step in this directory's Dockerfile — the
# offline cache lookup is revision-keyed, so a mismatch fails at runtime.
TOKENIZER_MODEL = "sentence-transformers/all-MiniLM-L6-v2"
# A public git commit hash, not a credential.
TOKENIZER_REVISION = "1110a243fdf4706b3f48f1d95db1a4f5529b4d41"  # pragma: allowlist secret


def pil_to_png(pil) -> bytes:
    buf = io.BytesIO()
    pil.save(buf, format="PNG")
    return buf.getvalue()


def parallel_map(fn: Callable, items: list, max_workers: int) -> list:
    """Map fn over items concurrently, preserving input order (latency-bound
    Bedrock calls)."""
    if not items:
        return []
    if max_workers <= 1 or len(items) == 1:
        return [fn(x) for x in items]
    with ThreadPoolExecutor(max_workers=max_workers) as ex:
        return list(ex.map(fn, items))


def build_converter(cfg: Settings = SETTINGS) -> DocumentConverter:
    # PDF: full page images + picture crops (needed for the multimodal embeds).
    pdf_opts = PdfPipelineOptions()
    pdf_opts.images_scale = cfg.images_scale
    pdf_opts.generate_page_images = True
    pdf_opts.generate_picture_images = True
    pdf_opts.do_ocr = False
    # In Lambda the filesystem is read-only outside /tmp; load models from the
    # image-baked artifacts path so Docling never fetches from HuggingFace.
    artifacts_path = os.environ.get("DOCLING_ARTIFACTS")
    if artifacts_path:
        pdf_opts.artifacts_path = artifacts_path

    pptx_opts = PaginatedPipelineOptions()
    pptx_opts.images_scale = cfg.images_scale
    pptx_opts.generate_picture_images = True

    xlsx_opts = PaginatedPipelineOptions()

    return DocumentConverter(
        allowed_formats=[InputFormat.PDF, InputFormat.PPTX, InputFormat.XLSX],
        format_options={
            InputFormat.PDF: PdfFormatOption(pipeline_options=pdf_opts),
            InputFormat.PPTX: PowerpointFormatOption(pipeline_options=pptx_opts),
            InputFormat.XLSX: ExcelFormatOption(pipeline_options=xlsx_opts),
        },
    )


def collect_page_text(doc, chunker: HybridChunker) -> dict[int, str]:
    by_page: dict[int, list[str]] = defaultdict(list)
    for chunk in chunker.chunk(doc):
        pages = set()
        for item in getattr(chunk.meta, "doc_items", []) or []:
            for prov in getattr(item, "prov", []) or []:
                if getattr(prov, "page_no", None) is not None:
                    pages.add(prov.page_no)
        for p in pages:
            by_page[p].append(chunk.text)
    return {p: "\n".join(texts) for p, texts in by_page.items()}


# ── PPTX slide rendering (best-effort; no LibreOffice in the image = skipped) ─
def render_pptx_slides(path: Path) -> dict[int, bytes]:
    out: dict[int, bytes] = {}
    try:
        with tempfile.TemporaryDirectory() as td:
            tdp = Path(td)
            subprocess.run(
                ["soffice", "--headless", "--convert-to", "pdf",
                 "--outdir", str(tdp), str(path)],
                check=True, capture_output=True, timeout=600,
            )
            pdf_path = tdp / (path.stem + ".pdf")
            if not pdf_path.exists():
                return {}
            subprocess.run(
                ["pdftoppm", "-png", "-r", "150", str(pdf_path), str(tdp / "slide")],
                check=True, capture_output=True, timeout=600,
            )
            for png_file in sorted(tdp.glob("slide-*.png")):
                num = int(png_file.stem.split("-")[-1])
                out[num] = png_file.read_bytes()
    except (subprocess.CalledProcessError, subprocess.TimeoutExpired,
            FileNotFoundError, ValueError) as e:
        logger.warning("PPTX slide render unavailable (%s); skipping page images", e)
    return out


# ── PPTX native charts (Docling drops these) ──────────────────────────────────
def _iter_chart_shapes(shapes):
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    for sh in shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _iter_chart_shapes(sh.shapes)
        elif getattr(sh, "has_chart", False):
            yield sh


def _fmt_num(v) -> str:
    if isinstance(v, (int, float)):
        if float(v).is_integer():
            return str(int(v))
        return f"{v:.2f}".rstrip("0").rstrip(".")
    return str(v)


def _chart_to_markdown(chart) -> tuple[Optional[str], str]:
    title = None
    if chart.has_title:
        title = (chart.chart_title.text_frame.text or "").strip() or None
    cats: list[str] = []
    if chart.plots:
        try:
            cats = [("" if c is None else str(c)) for c in chart.plots[0].categories]
        except (ValueError, AttributeError):
            cats = []
    series = []
    for s in chart.series:
        name = (s.name or "").strip() if s.name else ""
        vals = ["" if v is None else _fmt_num(v) for v in s.values]
        series.append((name, vals))
    lines = []
    if cats:
        lines.append("| series | " + " | ".join(cats) + " |")
        lines.append("| --- " * (len(cats) + 1) + "|")
        for name, vals in series:
            cells = (vals + [""] * len(cats))[: len(cats)]
            lines.append(f"| {name} | " + " | ".join(cells) + " |")
    else:
        lines += ["| series | values |", "| --- | --- |"]
        for name, vals in series:
            lines.append(f"| {name} | " + ", ".join(vals) + " |")
    return title, "\n".join(lines)


def extract_pptx_charts(path: Path) -> list[dict]:
    from pptx import Presentation
    out: list[dict] = []
    prs = Presentation(str(path))
    for slide_no, slide in enumerate(prs.slides, start=1):
        for sh in _iter_chart_shapes(slide.shapes):
            try:
                chart = sh.chart
                title, md = _chart_to_markdown(chart)
                if not md.strip():
                    continue
                try:
                    ctype = str(chart.chart_type)
                except (ValueError, AttributeError):
                    ctype = "unknown"
                out.append(dict(slide_no=slide_no, idx=len(out), title=title,
                                chart_type=ctype, markdown=md))
            except Exception as e:  # noqa: BLE001
                logger.warning("chart skipped on slide %d of %s (%s)",
                               slide_no, path.name, e)
    return out


@dataclass
class Ingestor:
    cfg: Settings
    bedrock: Bedrock
    converter: DocumentConverter
    chunker: HybridChunker = field(default_factory=HybridChunker)

    def _embed_batched(self, texts: list[str]) -> list[list[float]]:
        if not texts:
            return []
        size = max(1, self.cfg.embed_batch_size)
        slices = [texts[i:i + size] for i in range(0, len(texts), size)]
        results = parallel_map(self.bedrock.embed_texts, slices, self.cfg.max_workers)
        return [vec for batch in results for vec in batch]

    def _text_records(self, doc, source, document_id, ftype) -> list[dict]:
        chunks = [(i, c.text) for i, c in enumerate(self.chunker.chunk(doc))
                  if c.text.strip()]
        vectors = self._embed_batched([text for _, text in chunks])
        return [
            dict(id=f"{document_id}:text:{i}", document_id=document_id,
                 source_file=source, file_type=ftype, page_no=None,
                 modality="text", ref=None, caption=None,
                 text=text, vector=vec, image_bytes=None)
            for (i, text), vec in zip(chunks, vectors)
        ]

    def _figure_records(self, doc, source, document_id, ftype) -> list[dict]:
        figs = []
        for el, _level in doc.iterate_items():
            if not isinstance(el, PictureItem):
                continue
            pil = el.get_image(doc)
            if pil is None:
                continue
            png = pil_to_png(pil)
            figs.append(dict(
                self_ref=el.self_ref, png=png,
                caption=el.caption_text(doc=doc) or None,
                page_no=el.prov[0].page_no if el.prov else None,
            ))

        def build(f: dict) -> dict:
            if self.cfg.describe_figures:
                text = self.bedrock.describe(f["png"])
            else:
                text = f["caption"] or "figure"
            context = f"{source} | {f['caption'] or ''} | {text}".strip()
            return dict(
                id=f"{document_id}:fig:{f['self_ref']}", document_id=document_id,
                source_file=source, file_type=ftype, page_no=f["page_no"],
                modality="figure", ref=f["self_ref"], caption=f["caption"],
                text=text, vector=self.bedrock.embed_interleaved(context, f["png"]),
                image_bytes=f["png"],
            )

        return parallel_map(build, figs, self.cfg.max_workers)

    def _chart_records(self, path, source, document_id, ftype) -> list[dict]:
        charts = extract_pptx_charts(path)
        if not charts:
            return []

        def to_text(c: dict) -> str:
            header = " | ".join(x for x in (c["title"], f"chart: {c['chart_type']}") if x)
            return f"{header}\n{c['markdown']}" if header else c["markdown"]

        texts = [to_text(c) for c in charts]
        vectors = self._embed_batched(texts)
        return [
            dict(id=f"{document_id}:chart:{c['slide_no']}:{c['idx']}",
                 document_id=document_id, source_file=source, file_type=ftype,
                 page_no=c["slide_no"], modality="table", ref=None,
                 caption=c["title"] or f"{c['chart_type']} chart",
                 text=text, vector=vec, image_bytes=None)
            for c, text, vec in zip(charts, texts, vectors)
        ]

    def _table_records(self, doc, source, document_id, ftype) -> list[dict]:
        tables = []
        for el, _level in doc.iterate_items():
            if not isinstance(el, TableItem):
                continue
            try:
                md = el.export_to_markdown(doc=doc)
            except TypeError:
                md = el.export_to_markdown()
            if not md or not md.strip():
                continue
            tables.append(dict(
                self_ref=el.self_ref,
                page_no=el.prov[0].page_no if el.prov else None,
                caption=el.caption_text(doc=doc) or None, markdown=md,
            ))
        vectors = self._embed_batched([t["markdown"] for t in tables])
        return [
            dict(id=f"{document_id}:table:{t['self_ref']}", document_id=document_id,
                 source_file=source, file_type=ftype, page_no=t["page_no"],
                 modality="table", ref=t["self_ref"], caption=t["caption"],
                 text=t["markdown"], vector=vec, image_bytes=None)
            for t, vec in zip(tables, vectors)
        ]

    def _figures_block(self, figs: list[dict]) -> str:
        if not figs:
            return ""
        lines = ["", "--- Figures on this page ---"]
        for f in figs:
            cap = f" (caption: {f['caption']})" if f.get("caption") else ""
            lines.append(f"[figure id={f['id']}]{cap}")
            lines.append(f["text"])
        return "\n".join(lines)

    def _page_records(self, doc, source, document_id, ftype, page_text,
                      figs_by_page, slide_imgs=None) -> list[dict]:
        pages: list[tuple[int, Optional[bytes], str]] = []

        def collect(page_no: int, png: Optional[bytes], text: str):
            body = text or f"{source} page {page_no}"
            body = body + self._figures_block(figs_by_page.get(page_no, []))
            pages.append((page_no, png, body))

        if ftype == "pdf":
            for page_no, page in doc.pages.items():
                png = (pil_to_png(page.image.pil_image)
                       if page.image and page.image.pil_image else None)
                collect(page_no, png, page_text.get(page_no, ""))
        elif ftype == "pptx":
            slide_imgs = slide_imgs or {}
            page_nos = (set(page_text) | set(slide_imgs)
                        | set(figs_by_page) | set(doc.pages or {}))
            for page_no in sorted(p for p in page_nos if p is not None):
                collect(page_no, slide_imgs.get(page_no), page_text.get(page_no, ""))

        def build(item: tuple[int, Optional[bytes], str]) -> dict:
            page_no, png, body = item
            if png is not None:
                vec = self.bedrock.embed_interleaved(body, png)
            else:
                vec = self.bedrock.embed_text(body)
            return dict(
                id=f"{document_id}:page:{page_no}", document_id=document_id,
                source_file=source, file_type=ftype, page_no=page_no,
                modality="page", ref=None, caption=None, text=body,
                vector=vec, image_bytes=png,
            )

        return parallel_map(build, pages, self.cfg.max_workers)

    def ingest_file(self, path: Path, document_id: str,
                    source_name: Optional[str] = None) -> list[dict]:
        ext = path.suffix.lower()
        ftype = {".pdf": "pdf", ".pptx": "pptx", ".xlsx": "xlsx"}.get(ext)
        if ftype is None:
            raise ValueError(f"unsupported file type: {ext}")

        doc = self.converter.convert(path).document
        source = source_name or path.name
        page_text = collect_page_text(doc, self.chunker)

        records: list[dict] = []
        records += self._text_records(doc, source, document_id, ftype)
        records += self._table_records(doc, source, document_id, ftype)
        if ftype == "pptx":
            records += self._chart_records(path, source, document_id, ftype)

        fig_records: list[dict] = []
        if ftype in ("pdf", "pptx"):
            fig_records = self._figure_records(doc, source, document_id, ftype)
            records += fig_records

        figs_by_page: dict[int, list[dict]] = defaultdict(list)
        for f in fig_records:
            if f["page_no"] is not None:
                figs_by_page[f["page_no"]].append(f)

        slide_imgs = render_pptx_slides(path) if ftype == "pptx" else None
        records += self._page_records(doc, source, document_id, ftype, page_text,
                                      figs_by_page, slide_imgs)

        logger.info("  -> %d records (%s)", len(records), source)
        return records


def _build_chunker() -> HybridChunker:
    """HybridChunker with an explicit tokenizer. Its default tokenizer
    (all-MiniLM-L6-v2) can't auto-detect max_tokens under HF_HUB_OFFLINE — the
    HF config's model_max_length isn't in the local cache — so we construct it
    ourselves (weights baked into HF_HOME at image build; 256 = the model's
    context window)."""
    from docling_core.transforms.chunker.tokenizer.huggingface import (
        HuggingFaceTokenizer,
    )
    from transformers import AutoTokenizer

    # The revision IS pinned, to the immutable commit in TOKENIZER_REVISION
    # above; bandit only recognizes a literal here and can't follow the
    # module-level constant.
    tok = AutoTokenizer.from_pretrained(  # nosec B615
        TOKENIZER_MODEL, revision=TOKENIZER_REVISION
    )
    return HybridChunker(tokenizer=HuggingFaceTokenizer(tokenizer=tok, max_tokens=256))


def build_ingestor(cfg: Settings = SETTINGS, vlm_model: Optional[str] = None) -> Ingestor:
    if vlm_model and vlm_model != cfg.vlm_model:
        cfg = replace(cfg, vlm_model=vlm_model)
    return Ingestor(cfg=cfg, bedrock=Bedrock(cfg),
                    converter=build_converter(cfg), chunker=_build_chunker())


# ── LanceDB write + indexing ──────────────────────────────────────────────────
def get_table(cfg: Settings = SETTINGS):
    """Open (or create) the LanceDB table at the configured URI (S3 or local)."""
    db = lancedb.connect(cfg.db_uri)
    return db.create_table(cfg.table_name, schema=DocChunk, exist_ok=True)


def ensure_indexes(table, cfg: Settings = SETTINGS):
    try:
        table.create_fts_index(
            "text", base_tokenizer="ngram", ngram_min_length=2,
            ngram_max_length=3, lower_case=True, stem=False,
            remove_stop_words=False, replace=True,
        )
    except Exception as e:  # noqa: BLE001
        logger.warning("fts index not built (%s)", e)

    n = table.count_rows()
    if n >= cfg.ann_min_rows:
        try:
            table.create_index(metric="cosine", vector_column_name="vector")
            logger.info("vector ANN index built (%d rows)", n)
        except Exception as e:  # noqa: BLE001
            logger.warning("vector index not built (%s)", e)
    else:
        logger.info("skipping vector ANN index (%d rows < %d)", n, cfg.ann_min_rows)
