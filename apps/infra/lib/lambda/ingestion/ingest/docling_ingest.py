"""Docling document ingestion for the graph-ontology experiment.

Adapted from apps/worker-semantic-ingestion/ingest.py. Docling remains the
backbone, with three recovery techniques ported from the worker pipeline for
content Docling misses:

* native PPTX charts — Docling's pptx backend drops chart objects
  (docling#2306); python-pptx exposes each chart's cached categories/series,
  which we render as markdown tables into the slide's text.
* PPTX speaker notes — Docling parses them into ContentLayer.NOTES, which
  HybridChunker does not traverse; read them via python-pptx instead.
* VLM fallback for image-only pages — pages whose extracted text is below
  VLM_PAGE_MIN_CHARS are rendered (PPTX: LibreOffice -> pdftoppm, PDF:
  pypdfium2) and transcribed by the Bedrock model (describe_image). This is
  what rescues screenshot-heavy decks that otherwise yield 0 infons.
  Best-effort: if the toolchain or credentials are missing, ingestion
  proceeds with whatever text Docling found.

Converted text is cached as JSON under .docling-cache/ keyed by
CACHE_VERSION + filename + mtime, so re-runs skip conversion (and VLM calls).
"""

from __future__ import annotations

import hashlib
import json
import logging
import os
import subprocess
import tempfile
from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path

from docling.chunking import HybridChunker
from docling.datamodel.base_models import InputFormat
from docling.datamodel.pipeline_options import (
    PaginatedPipelineOptions,
    PdfPipelineOptions,
)
from docling.document_converter import (
    DocumentConverter,
    ExcelFormatOption,
    PdfFormatOption,
    PowerpointFormatOption,
    WordFormatOption,
)

from bedrock import describe_image

log = logging.getLogger(__name__)

SUPPORTED_SUFFIXES = {".pdf", ".docx", ".pptx", ".xlsx"}
# In Lambda only /tmp is writable; DOCLING_CACHE_DIR points there. Falls back
# to a sibling dir for local runs.
CACHE_DIR = Path(
    os.environ.get("DOCLING_CACHE_DIR", str(Path(__file__).with_name(".docling-cache")))
)
# Bump whenever extraction logic changes so stale cache entries are ignored.
CACHE_VERSION = "v2"

# VLM fallback tuning: pages with less than MIN_CHARS of extracted text get
# rendered + described; per-document page cap bounds cost on huge documents.
VLM_PAGE_MIN_CHARS = 200
VLM_MAX_PAGES_PER_DOC = 40
RENDER_DPI = 110  # slide/page render resolution for the VLM


@dataclass(frozen=True)
class IngestedDocument:
    document_id: str  # stable id derived from the file name
    title: str
    source_path: str
    text: str  # full document text (page texts joined)
    content_hash: str  # sha256 of the text — feeds infon ids
    page_count: int
    # Extraction provenance (for validation/reporting).
    docling_chars: int = 0
    chart_count: int = 0
    notes_chars: int = 0
    vlm_pages: int = 0
    vlm_chars: int = 0


def build_converter() -> DocumentConverter:
    """Text-only converter: no OCR, no page/picture image generation."""
    pdf_opts = PdfPipelineOptions()
    pdf_opts.do_ocr = False
    pdf_opts.generate_page_images = False
    pdf_opts.generate_picture_images = False

    # In Lambda the filesystem is read-only except /tmp, so Docling must NOT try
    # to fetch its layout model from HuggingFace at runtime. DOCLING_ARTIFACTS
    # points at the models baked into the image at build time; setting
    # artifacts_path makes Docling load from disk and skip all network/HF-cache
    # writes. Local runs (env unset) keep Docling's default download behavior.
    artifacts_path = os.environ.get("DOCLING_ARTIFACTS")
    if artifacts_path:
        pdf_opts.artifacts_path = artifacts_path

    paginated_opts = PaginatedPipelineOptions()

    return DocumentConverter(
        allowed_formats=[
            InputFormat.PDF,
            InputFormat.DOCX,
            InputFormat.PPTX,
            InputFormat.XLSX,
        ],
        format_options={
            InputFormat.PDF: PdfFormatOption(pipeline_options=pdf_opts),
            InputFormat.DOCX: WordFormatOption(pipeline_options=paginated_opts),
            InputFormat.PPTX: PowerpointFormatOption(pipeline_options=paginated_opts),
            InputFormat.XLSX: ExcelFormatOption(pipeline_options=paginated_opts),
        },
    )


def collect_page_text(doc, chunker: HybridChunker) -> dict[int, str]:
    """Map page_no -> concatenated chunk text on that page."""
    by_page: dict[int, list[str]] = defaultdict(list)
    for chunk in chunker.chunk(doc):
        pages = set()
        for item in getattr(chunk.meta, "doc_items", []) or []:
            for prov in getattr(item, "prov", []) or []:
                if getattr(prov, "page_no", None) is not None:
                    pages.add(prov.page_no)
        if not pages:
            pages = {0}  # non-paginated formats (xlsx/docx without prov)
        for p in pages:
            by_page[p].append(chunk.text)
    return {p: "\n".join(texts) for p, texts in by_page.items()}


# ── PPTX chart recovery (ported from worker-semantic-ingestion) ──────────────


def _iter_chart_shapes(shapes):
    """Yield chart-bearing shapes, descending into group shapes."""
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    for sh in shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _iter_chart_shapes(sh.shapes)
        elif getattr(sh, "has_chart", False):
            yield sh


def _fmt_num(v) -> str:
    if isinstance(v, (int, float)):
        if float(v).is_integer():
            return str(int(v))
        return f"{v:.2f}".rstrip("0").rstrip(".")
    return str(v)


def _chart_to_markdown(chart) -> str:
    """Render a python-pptx chart as a markdown table (categories as columns,
    one row per series); series-only listing when there are no categories."""
    title = None
    if chart.has_title:
        title = (chart.chart_title.text_frame.text or "").strip() or None

    cats: list[str] = []
    if chart.plots:
        try:
            cats = [("" if c is None else str(c)) for c in chart.plots[0].categories]
        except (ValueError, AttributeError):
            cats = []

    series = []
    for s in chart.series:
        name = (s.name or "").strip() if s.name else ""
        vals = ["" if v is None else _fmt_num(v) for v in s.values]
        series.append((name, vals))

    lines = [f"Chart: {title}" if title else "Chart:"]
    if cats:
        lines.append("| series | " + " | ".join(cats) + " |")
        lines.append("| --- " * (len(cats) + 1) + "|")
        for name, vals in series:
            cells = (vals + [""] * len(cats))[: len(cats)]
            lines.append(f"| {name} | " + " | ".join(cells) + " |")
    else:
        lines += ["| series | values |", "| --- | --- |"]
        for name, vals in series:
            lines.append(f"| {name} | " + ", ".join(vals) + " |")
    return "\n".join(lines)


def extract_pptx_charts(path: Path) -> dict[int, list[str]]:
    """slide_no (1-based) -> markdown tables for each chart on that slide.

    Best-effort: a malformed chart is skipped, never fails the file.
    """
    from pptx import Presentation

    out: dict[int, list[str]] = defaultdict(list)
    try:
        prs = Presentation(str(path))
    except Exception as e:  # noqa: BLE001
        log.warning("python-pptx failed to open %s (%s)", path.name, e)
        return {}
    for slide_no, slide in enumerate(prs.slides, start=1):
        for sh in _iter_chart_shapes(slide.shapes):
            try:
                md = _chart_to_markdown(sh.chart)
                if md.strip():
                    out[slide_no].append(md)
            except Exception as e:  # noqa: BLE001
                log.warning(
                    "chart skipped on slide %d of %s (%s)", slide_no, path.name, e
                )
    return dict(out)


def extract_pptx_notes(path: Path) -> dict[int, str]:
    """slide_no (1-based) -> speaker-notes text (Docling parses notes into
    ContentLayer.NOTES, which the chunker skips — so read them directly)."""
    from pptx import Presentation

    out: dict[int, str] = {}
    try:
        prs = Presentation(str(path))
    except Exception:  # noqa: BLE001 - already warned in chart extraction
        return {}
    for slide_no, slide in enumerate(prs.slides, start=1):
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            text = slide.notes_slide.notes_text_frame.text.strip()
            if text:
                out[slide_no] = text
    return out


# ── Page rendering for the VLM fallback ──────────────────────────────────────


def render_pptx_slides(path: Path) -> dict[int, bytes]:
    """Render each slide to PNG via headless LibreOffice -> pdftoppm.

    Returns {slide_number (1-based): png_bytes}; {} if the toolchain is
    missing so ingestion proceeds without the VLM fallback.
    """
    out: dict[int, bytes] = {}
    try:
        with tempfile.TemporaryDirectory() as td:
            tdp = Path(td)
            subprocess.run(
                ["soffice", "--headless", "--convert-to", "pdf",
                 "--outdir", str(tdp), str(path)],
                check=True, capture_output=True, timeout=600,
            )
            pdf_path = tdp / (path.stem + ".pdf")
            if not pdf_path.exists():
                return {}
            subprocess.run(
                ["pdftoppm", "-png", "-r", str(RENDER_DPI),
                 str(pdf_path), str(tdp / "slide")],
                check=True, capture_output=True, timeout=600,
            )
            for png_file in sorted(tdp.glob("slide-*.png")):
                num = int(png_file.stem.split("-")[-1])
                out[num] = png_file.read_bytes()
    except (subprocess.CalledProcessError, subprocess.TimeoutExpired,
            FileNotFoundError, ValueError) as e:
        log.warning("PPTX slide render unavailable for %s (%s)", path.name, e)
    return out


def render_pdf_pages(path: Path, page_nos: list[int]) -> dict[int, bytes]:
    """Render selected PDF pages (1-based) to PNG via pypdfium2."""
    import io

    import pypdfium2 as pdfium

    out: dict[int, bytes] = {}
    try:
        pdf = pdfium.PdfDocument(str(path))
        for page_no in page_nos:
            if not 1 <= page_no <= len(pdf):
                continue
            bitmap = pdf[page_no - 1].render(scale=RENDER_DPI / 72)
            buf = io.BytesIO()
            bitmap.to_pil().save(buf, format="PNG")
            out[page_no] = buf.getvalue()
    except Exception as e:  # noqa: BLE001
        log.warning("PDF page render failed for %s (%s)", path.name, e)
    return out


def _true_page_count(path: Path, suffix: str) -> int:
    """Real page/slide count from the source file — Docling's page_text only
    covers pages where it found content, so fully image-only pages are absent
    from it entirely (and they are exactly the VLM fallback's targets)."""
    try:
        if suffix == ".pptx":
            from pptx import Presentation

            return len(Presentation(str(path)).slides)
        if suffix == ".pdf":
            import pypdfium2 as pdfium

            return len(pdfium.PdfDocument(str(path)))
    except Exception as e:  # noqa: BLE001
        log.warning("page count failed for %s (%s)", path.name, e)
    return 0


def _vlm_describe_pages(
    path: Path, suffix: str, page_text: dict[int, str], title: str
) -> dict[int, str]:
    """Describe low-text pages with the VLM; page_no -> transcription."""
    n_pages = _true_page_count(path, suffix)
    if n_pages == 0:
        return {}
    low = sorted(
        p for p in range(1, n_pages + 1)
        if len(page_text.get(p, "").strip()) < VLM_PAGE_MIN_CHARS
    )
    if not low:
        return {}
    if len(low) > VLM_MAX_PAGES_PER_DOC:
        log.warning(
            "%s: %d low-text pages, describing first %d only",
            path.name, len(low), VLM_MAX_PAGES_PER_DOC,
        )
        low = low[:VLM_MAX_PAGES_PER_DOC]

    if suffix == ".pptx":
        renders = render_pptx_slides(path)
        renders = {p: png for p, png in renders.items() if p in set(low)}
    elif suffix == ".pdf":
        renders = render_pdf_pages(path, low)
    else:
        return {}

    out: dict[int, str] = {}
    for page_no in low:
        png = renders.get(page_no)
        if png is None:
            continue
        text = describe_image(png, context=title)
        if text:
            out[page_no] = text
    log.info(
        "%s: VLM described %d/%d low-text page(s)", path.name, len(out), len(low)
    )
    return out


# ── Ingestion ────────────────────────────────────────────────────────────────


def _document_id(path: Path) -> str:
    """Stable, filename-derived id (re-runs of the same file MERGE cleanly)."""
    return "doc-" + hashlib.sha256(path.name.encode("utf-8")).hexdigest()[:12]


def _cache_path(path: Path) -> Path:
    key = f"{CACHE_VERSION}:{path.name}:{path.stat().st_mtime_ns}"
    return CACHE_DIR / (hashlib.sha256(key.encode()).hexdigest()[:16] + ".json")


def ingest_file(
    path: Path, converter: DocumentConverter, chunker: HybridChunker
) -> IngestedDocument:
    """Convert one file to text (cached), returning an IngestedDocument."""
    cache = _cache_path(path)
    if cache.exists():
        data = json.loads(cache.read_text(encoding="utf-8"))
        log.info("cache hit: %s", path.name)
    else:
        log.info("converting: %s", path.name)
        doc = converter.convert(path).document
        page_text = collect_page_text(doc, chunker)
        docling_chars = sum(len(t) for t in page_text.values())
        suffix = path.suffix.lower()

        chart_count = 0
        notes_chars = 0
        if suffix == ".pptx":
            for slide_no, charts in extract_pptx_charts(path).items():
                block = "\n\n".join(charts)
                chart_count += len(charts)
                page_text[slide_no] = (
                    page_text.get(slide_no, "") + "\n\n" + block
                ).strip()
            for slide_no, notes in extract_pptx_notes(path).items():
                notes_chars += len(notes)
                page_text[slide_no] = (
                    page_text.get(slide_no, "") + "\n\nNotes: " + notes
                ).strip()

        vlm_text = _vlm_describe_pages(path, suffix, page_text, path.stem)
        for page_no, text in vlm_text.items():
            page_text[page_no] = (
                page_text.get(page_no, "") + "\n\n" + text
            ).strip()

        text = "\n\n".join(page_text[p] for p in sorted(page_text))
        data = {
            "text": text,
            "page_count": len(page_text),
            "docling_chars": docling_chars,
            "chart_count": chart_count,
            "notes_chars": notes_chars,
            "vlm_pages": len(vlm_text),
            "vlm_chars": sum(len(t) for t in vlm_text.values()),
        }
        CACHE_DIR.mkdir(exist_ok=True)
        cache.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")

    text = data["text"]
    return IngestedDocument(
        document_id=_document_id(path),
        title=path.stem,
        source_path=str(path),
        text=text,
        content_hash=hashlib.sha256(text.encode("utf-8")).hexdigest(),
        page_count=data["page_count"],
        docling_chars=data.get("docling_chars", 0),
        chart_count=data.get("chart_count", 0),
        notes_chars=data.get("notes_chars", 0),
        vlm_pages=data.get("vlm_pages", 0),
        vlm_chars=data.get("vlm_chars", 0),
    )


def scan_docs(docs_dir: Path) -> list[Path]:
    return sorted(
        p
        for p in docs_dir.iterdir()
        if p.is_file()
        and p.suffix.lower() in SUPPORTED_SUFFIXES
        # Skip Office lock files (~$foo.pptx) and hidden files.
        and not p.name.startswith(("~$", "."))
    )
